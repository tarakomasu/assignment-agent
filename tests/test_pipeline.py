import copy
import json

import pytest
from docx import Document
from pypdf import PdfWriter

from assignment_agent.models import AgentError, parse_plan, parse_sources
from assignment_agent.pipeline import review_result, run
from test_runner import SOURCE

PLAN = {"lecture_title": "標準入力", "missing_information": [], "assignments": [
    {"id": "task_01", "title": "携帯電話の料金", "page": 1,
     "requirements": "基本3780円、通話20円/分、メール5円/通",
     "tests": [{"stdin": "10\n20\n", "expected_stdout": "4080\n"},
               {"stdin": "0\n0\n", "expected_stdout": "3780\n"}]}]}


@pytest.fixture
def pdf(tmp_path):
    path = tmp_path / "講義 資料.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    with path.open("wb") as f:
        writer.write(f)
    return path


def fake_ai(responses):
    class Fake:
        def __init__(self, *args, **kwargs):
            self.responses = iter(copy.deepcopy(responses))

        def ask(self, prompt):
            return next(self.responses)
    return Fake


def generate(source=SOURCE, task_id="task_01"):
    return {"solutions": [{"id": task_id, "source": source}]}


def test_real_c_png_word_flow(pdf, tmp_path):
    result = run(pdf, tmp_path / "out", tmp_path / "work",
                 ai_factory=fake_ai([PLAN, generate(), {"valid": True, "issues": []}]))
    document = Document(result)
    text = "\n".join(p.text for p in document.paragraphs)
    assert "携帯電話の料金" in text
    assert "scanf" in text
    assert document.inline_shapes
    assert all(shape.width < document.sections[0].page_width - document.sections[0].left_margin - document.sections[0].right_margin
               for shape in document.inline_shapes)
    tests = json.loads((result.parent / "task_01/tests.json").read_text(encoding="utf-8"))
    assert all(t["passed"] for t in tests)
    assert tests[0]["stdout"].replace("\r\n", "\n") == "4080\n"


def test_wrong_code_repaired_without_changing_expectations(pdf, tmp_path, monkeypatch):
    # One artifact pass is already covered above; keep this check focused on retries.
    monkeypatch.setattr("assignment_agent.pipeline.render_result", lambda *a: [])
    result = run(pdf, tmp_path / "out", tmp_path / "work", ai_factory=fake_ai([
        PLAN, generate(SOURCE.replace("3780+", "1000+")), generate(), {"valid": True, "issues": []}]))
    tests = json.loads((result.parent / "task_01/tests.json").read_text(encoding="utf-8"))
    assert tests[0]["expected_stdout"] == "4080\n"
    assert tests[0]["passed"]


def test_failure_never_publishes_word(pdf, tmp_path):
    bad = generate(SOURCE.replace("3780+", "1000+"))
    with pytest.raises(AgentError):
        run(pdf, tmp_path / "out", tmp_path / "work", ai_factory=fake_ai([PLAN, bad, bad, bad]))
    assert not list(tmp_path.glob("out/**/*.docx"))
    assert list(tmp_path.glob("work/*/failure.json"))


def test_review_failure_never_publishes_word(pdf, tmp_path):
    rejected = {"valid": False, "issues": ["今回の課題が不足"]}
    with pytest.raises(AgentError):
        run(pdf, tmp_path / "out", tmp_path / "work", ai_factory=fake_ai([
            PLAN, generate(), rejected, generate(), rejected, generate(), rejected]))
    assert not list(tmp_path.glob("out/**/*.docx"))


def test_multiple_assignments_all_processed(pdf, tmp_path, monkeypatch):
    monkeypatch.setattr("assignment_agent.pipeline.render_result", lambda *a: [])
    plan = copy.deepcopy(PLAN)
    second = copy.deepcopy(plan["assignments"][0])
    second.update(id="task_02", title="課題2")
    plan["assignments"].append(second)
    sources = generate()
    sources["solutions"].append({"id": "task_02", "source": SOURCE})
    result = run(pdf, tmp_path / "out", tmp_path / "work", ai_factory=fake_ai([plan, sources, {"valid": True, "issues": []}]))
    assert (result.parent / "task_02/solution.c").is_file()
    text = "\n".join(p.text for p in Document(result).paragraphs)
    assert "携帯電話の料金" in text and "課題2" in text


@pytest.mark.parametrize("mutator", [
    lambda p: p.update(missing_information=["氏名が必要です"]),
    lambda p: p["assignments"][0].update(id="../escape"),
    lambda p: p["assignments"][0].update(page=2),
    lambda p: p["assignments"][0].update(tests=[]),
    lambda p: p.update(assignments=[]),
    lambda p: p.update(missing_information="wrong type"),
])
def test_invalid_or_incomplete_plan(mutator):
    plan = copy.deepcopy(PLAN)
    mutator(plan)
    with pytest.raises(AgentError):
        parse_plan(plan, 1)


def test_missing_solution_fails():
    _, tasks = parse_plan(PLAN, 1)
    with pytest.raises(AgentError):
        parse_sources({"solutions": []}, tasks)


@pytest.mark.parametrize("value", [{"valid": "true", "issues": []}, {"valid": True, "issues": ["bad"]}, {"valid": False, "issues": []}])
def test_invalid_review(value):
    with pytest.raises(AgentError):
        review_result(value)


def test_broken_pdf_stops_before_ai(tmp_path):
    path = tmp_path / "broken.pdf"
    path.write_text("not a PDF")
    with pytest.raises(AgentError):
        run(path, tmp_path / "out", tmp_path / "work", ai_factory=fake_ai([]))


def test_profile_stays_local_and_final_code_retested(pdf, tmp_path, monkeypatch):
    monkeypatch.setattr("assignment_agent.pipeline.render_result", lambda *a: [])
    plan = copy.deepcopy(PLAN)
    plan["assignments"][0]["tests"] = [{"stdin": "", "expected_stdout": "__STUDENT_NAME__ / __STUDENT_ID__\n"}]
    source = '#include <stdio.h>\nint main(void) { printf("__STUDENT_NAME__ / __STUDENT_ID__\\n"); return 0; }'
    profile = {"name": '学生"太郎', "student_id": "12345678"}
    base = fake_ai([plan, generate(source), {"valid": True, "issues": []}])

    class PrivateAI(base):
        def ask(self, prompt):
            assert profile["name"] not in prompt
            assert profile["student_id"] not in prompt
            return super().ask(prompt)

    result = run(pdf, tmp_path / "out", tmp_path / "work", ai_factory=PrivateAI, profile=profile)
    tests = json.loads((result.parent / "task_01/tests.json").read_text(encoding="utf-8"))
    assert tests[0]["stdout"].replace("\r\n", "\n") == '学生"太郎 / 12345678\n'
    assert tests[0]["passed"]
    text = "\n".join(p.text for p in Document(result).paragraphs)
    assert profile["name"] in text and profile["student_id"] in text

import html
import os
import textwrap
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.image.image import Image
from playwright.sync_api import Error, sync_playwright

from .gemini import app_root
from .models import AgentError


def render_result(source, tests, folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    # Split long code/output into bounded panels to keep screenshots readable in Word.
    sections = [("Cプログラム", source)]
    for i, test in enumerate(tests, 1):
        sections.append((f"テスト {i}", "入力\n" + (test["stdin"] or "（入力なし）\n")
                         + "\n実際の標準出力\n" + test["stdout"]))
    panels = []
    for title, content in sections:
        lines = [part for line in (content.expandtabs(4).splitlines() or [""])
                 for part in (textwrap.wrap(line, width=85, expand_tabs=False,
                                             replace_whitespace=False, drop_whitespace=False) or [""])]
        for start in range(0, len(lines), 24):
            panels.append((title + (" 続き" if start else ""), "\n".join(lines[start:start + 24])))
    browser_root = app_root() / "runtime" / "browsers"
    if browser_root.is_dir():
        os.environ["PLAYWRIGHT_BROWSERS_PATH"] = str(browser_root)
    paths = []
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch()
            try:
                page = browser.new_page(viewport={"width": 1100, "height": 800}, device_scale_factor=2)
                for i, (title, content) in enumerate(panels, 1):
                    markup = f"""<!doctype html><meta charset="utf-8"><style>
                    *{{box-sizing:border-box}}body{{margin:0;padding:28px;background:white;color:#111;
                    font-family:'Yu Gothic','Hiragino Sans',sans-serif}}article{{border:1px solid #aaa;padding:22px}}
                    h1{{font-size:21px;margin:0 0 6px}}small{{font-size:15px;color:#555}}
                    pre{{font:20px/1.5 Consolas,'Menlo','Yu Gothic',monospace;white-space:pre-wrap;
                    overflow-wrap:anywhere;margin:18px 0 0}}</style>
                    <article><h1>{html.escape(title)}</h1><small>ローカルCプログラムの実行記録</small>
                    <pre>{html.escape(content)}</pre></article>"""
                    (folder / f"result-{i:02d}.html").write_text(markup, encoding="utf-8")
                    page.set_content(markup)
                    page.evaluate("document.fonts.ready")
                    path = folder / f"result-{i:02d}.png"
                    page.screenshot(path=str(path), full_page=True)
                    paths.append(path)
            finally:
                browser.close()
    except Error as exc:
        raise AgentError("実行結果画像を作成できません。Windows用ZIPを完全に展開して再実行してください。") from exc
    return paths


def make_docx(title, assignments, images, destination, profile=None):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(20)
    section.left_margin = section.right_margin = Mm(20)
    normal = document.styles["Normal"]
    normal.font.name = "Yu Gothic"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Yu Gothic")
    for name in ("Title", "Heading 1", "Heading 2"):
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    document.add_paragraph(title, "Title")
    if profile:
        document.add_paragraph("氏名：" + profile["name"])
        document.add_paragraph("学籍番号：" + profile["student_id"])
    for index, task in enumerate(assignments):
        if index:
            document.add_page_break()
        document.add_heading(task.title, 1)
        document.add_heading("プログラム", 2)
        for line in task.source.splitlines():
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Yu Gothic")
        document.add_heading("実行結果", 2)
        for path in images[task.id]:
            # Each panel <= 24 lines; fit the available A4 width and height.
            image = Image.from_file(str(path))
            width = min(165, 205 * image.px_width / image.px_height)
            document.add_picture(str(path), width=Mm(width))
            drawing = document.paragraphs[-1]._p.xpath(".//wp:docPr")
            if drawing:
                drawing[0].set("descr", "ローカルでコンパイル・実行したCプログラムの検証記録")
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.title = title
    document.save(destination)
